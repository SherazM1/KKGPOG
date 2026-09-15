from __future__ import annotations

import io
from typing import Any

import openpyxl
from PIL import Image

from app.holiday_planograms.models import BosPlanogramProduct, ReferenceImage
from app.holiday_planograms.normalization import canonical_header, digits_only, int_value, text_value


QP_SHEET_HINTS = ("qtrpallet", "quarterpallet")

_TABLE_ALIASES: dict[str, tuple[str, ...]] = {
    "upc": ("upc",),
    "item_number": ("id", "item number", "wm item number"),
    "name": ("name", "product name"),
    "facings": ("facings", "facing"),
    "cpp": ("cpp",),
    "capacity": ("capacity",),
}


def _coerce_workbook_bytes(source_file: Any) -> bytes:
    if isinstance(source_file, bytes):
        return source_file
    if isinstance(source_file, bytearray):
        return bytes(source_file)
    if isinstance(source_file, str):
        with open(source_file, "rb") as handle:
            return handle.read()
    if hasattr(source_file, "getvalue"):
        return bytes(source_file.getvalue())
    if hasattr(source_file, "read"):
        data = source_file.read()
        if hasattr(source_file, "seek"):
            source_file.seek(0)
        return bytes(data)
    raise TypeError("Unsupported BOS workbook source.")


def select_planogram_sheet(sheet_names: list[str], hints: tuple[str, ...] = QP_SHEET_HINTS) -> str:
    scored: list[tuple[int, str]] = []
    for name in sheet_names:
        compact = canonical_header(name)
        score = sum(1 for hint in hints if hint in compact)
        if score:
            scored.append((score, name))
    if not scored:
        raise ValueError("Could not find a Quarter Pallet worksheet.")
    scored.sort(key=lambda item: (-item[0], item[1]))
    if len(scored) > 1 and scored[0][0] == scored[1][0]:
        tied = ", ".join(name for score, name in scored if score == scored[0][0])
        raise ValueError(f"Multiple equally strong Quarter Pallet worksheets found: {tied}.")
    return scored[0][1]


def _header_mapping(values: list[Any]) -> dict[str, int]:
    canonical = {canonical_header(value): index for index, value in enumerate(values) if text_value(value)}
    mapping: dict[str, int] = {}
    for logical, aliases in _TABLE_ALIASES.items():
        for alias in aliases:
            found = canonical.get(canonical_header(alias))
            if found is not None:
                mapping[logical] = found
                break
    return mapping


def _extract_table(ws: Any) -> tuple[list[BosPlanogramProduct], list[str]]:
    warnings: list[str] = []
    header_row: int | None = None
    mapping: dict[str, int] = {}
    for row_index in range(1, min(ws.max_row, 100) + 1):
        values = [ws.cell(row_index, col).value for col in range(1, ws.max_column + 1)]
        candidate = _header_mapping(values)
        if all(key in candidate for key in ("upc", "item_number", "name", "facings")):
            header_row = row_index
            mapping = candidate
            break
    if header_row is None:
        raise ValueError("Found QP worksheet but no structured UPC/ID/Name/Facings table.")

    products: list[BosPlanogramProduct] = []
    seen_ids: set[str] = set()
    seen_upcs: set[str] = set()
    blank_streak = 0
    for row_index in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(row_index, col).value for col in range(1, ws.max_column + 1)]
        raw = {
            key: values[column_index] if column_index < len(values) else None
            for key, column_index in mapping.items()
        }
        if not any(text_value(value) for value in raw.values()):
            blank_streak += 1
            if blank_streak >= 2:
                break
            continue
        blank_streak = 0
        facings = int_value(raw.get("facings"), allow_zero=False)
        if facings is None:
            warnings.append(f"BOS row {row_index}: malformed facings '{raw.get('facings')}'.")
            continue
        upc = digits_only(raw.get("upc"))
        item_number = digits_only(raw.get("item_number"))
        name = text_value(raw.get("name"))
        if not name:
            warnings.append(f"BOS row {row_index}: blank product name.")
        if item_number and item_number in seen_ids:
            warnings.append(f"BOS row {row_index}: duplicate ID {item_number}.")
        if upc and upc in seen_upcs:
            warnings.append(f"BOS row {row_index}: duplicate UPC {upc}.")
        seen_ids.add(item_number)
        seen_upcs.add(upc)
        products.append(
            BosPlanogramProduct(
                upc=upc,
                item_number=item_number,
                name=name,
                facings=facings,
                cpp=int_value(raw.get("cpp")),
                capacity=int_value(raw.get("capacity")),
                source_row=row_index,
                raw=raw,
            )
        )
    return products, warnings


def _extract_reference_images(ws: Any) -> list[ReferenceImage]:
    images: list[ReferenceImage] = []
    for index, image in enumerate(getattr(ws, "_images", []) or [], start=1):
        data = image._data()
        with Image.open(io.BytesIO(data)) as pil_image:
            width, height = pil_image.size
        anchor = getattr(image, "anchor", None)
        marker = getattr(anchor, "_from", None)
        row = int(getattr(marker, "row", 0) or 0)
        col = int(getattr(marker, "col", 0) or 0)
        images.append(
            ReferenceImage(
                name=f"worksheet_image_{index}",
                bytes_data=data,
                width=width,
                height=height,
                anchor_row=row,
                anchor_col=col,
            )
        )
    return images


def _select_reference_image(images: list[ReferenceImage]) -> tuple[ReferenceImage | None, list[str]]:
    warnings: list[str] = []
    if not images:
        return None, ["No embedded planogram reference images were found on the selected sheet."]
    plausible = [image for image in images if image.width >= 300 and image.height >= 250]
    if len(plausible) == 1:
        return plausible[0], warnings
    if len(plausible) > 1:
        largest_area = max(image.width * image.height for image in plausible)
        largest = [image for image in plausible if image.width * image.height == largest_area]
        if len(largest) == 1 and largest_area >= 1.5 * sorted((i.width * i.height for i in plausible), reverse=True)[1]:
            return largest[0], warnings
        names = ", ".join(f"{image.name}({image.width}x{image.height})" for image in plausible)
        return None, [f"Found {len(plausible)} candidate planogram images; automatic selection is ambiguous: {names}."]
    return max(images, key=lambda image: image.width * image.height), warnings


def load_bos_qp(source_file: Any) -> tuple[str, list[BosPlanogramProduct], ReferenceImage | None, list[str]]:
    payload = _coerce_workbook_bytes(source_file)
    wb = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
    try:
        sheet_name = select_planogram_sheet(wb.sheetnames)
        ws = wb[sheet_name]
        products, table_warnings = _extract_table(ws)
        images = _extract_reference_images(ws)
        reference_image, image_warnings = _select_reference_image(images)
        return sheet_name, products, reference_image, table_warnings + image_warnings
    finally:
        wb.close()
