from __future__ import annotations

import io
from typing import Any

import openpyxl

from app.holiday_planograms.extract_bos import _coerce_workbook_bytes, _extract_reference_images
from app.holiday_planograms.models import HolidayNreReference, ReferenceImage
from app.holiday_planograms.normalization import canonical_header, text_value


NRE_CONFIGURATIONS = ("4x78", "3x78", "3x60", "4x60")
NRE_SHEET_HINTS = ("2026nreholiday", "nreholiday")


def normalize_nre_configuration(configuration: str) -> str:
    compact = canonical_header(configuration).replace("by", "x")
    aliases = {
        "4x78": "4x78",
        "478": "4x78",
        "3x78": "3x78",
        "378": "3x78",
        "3x60": "3x60",
        "360": "3x60",
        "4x60": "4x60",
        "460": "4x60",
    }
    if compact not in aliases:
        raise ValueError(f"Unsupported Holiday NRE configuration: {configuration}.")
    return aliases[compact]


def select_nre_sheet(sheet_names: list[str]) -> str:
    scored: list[tuple[int, str]] = []
    for name in sheet_names:
        compact = canonical_header(name)
        score = sum(1 for hint in NRE_SHEET_HINTS if hint in compact)
        if score:
            scored.append((score, name))
    if not scored:
        raise ValueError("Could not find the Holiday NRE worksheet.")
    scored.sort(key=lambda item: (-item[0], item[1]))
    if len(scored) > 1 and scored[0][0] == scored[1][0]:
        tied = ", ".join(name for score, name in scored if score == scored[0][0])
        raise ValueError(f"Multiple equally strong Holiday NRE worksheets found: {tied}.")
    return scored[0][1]


def _title_token(configuration: str) -> str:
    return canonical_header(f"HOLIDAY NRE 2026 {configuration}")


def _find_nre_titles(ws: Any) -> dict[str, tuple[str, int, int]]:
    titles: dict[str, tuple[str, int, int]] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = text_value(cell.value)
            if not value:
                continue
            compact = canonical_header(value)
            for configuration in NRE_CONFIGURATIONS:
                if _title_token(configuration) in compact or (
                    "holidaynre2026" in compact and canonical_header(configuration) in compact
                ):
                    titles[configuration] = (value, int(cell.row), int(cell.column))
    return titles


def _select_image_near_title(
    configuration: str,
    title: tuple[str, int, int],
    images: list[ReferenceImage],
) -> tuple[ReferenceImage | None, list[str]]:
    warnings: list[str] = []
    if not images:
        return None, ["No embedded Holiday NRE reference images were found on the selected sheet."]

    title_text, title_row, title_col = title
    plausible = [image for image in images if image.width >= 250 and image.height >= 250]
    if not plausible:
        plausible = images

    def score(image: ReferenceImage) -> tuple[int, int, int]:
        image_row = image.anchor_row + 1
        image_col = image.anchor_col + 1
        below_or_near_title = 0 if image_row >= title_row else 1
        col_distance = abs(image_col - title_col)
        row_distance = abs(image_row - title_row)
        return below_or_near_title, col_distance, row_distance

    ranked = sorted(plausible, key=score)
    if len(ranked) > 1 and score(ranked[0]) == score(ranked[1]):
        tied = ", ".join(f"{image.name}({image.width}x{image.height})" for image in ranked[:2])
        warnings.append(
            f"Holiday NRE {configuration} reference selection is ambiguous near title '{title_text}': {tied}."
        )
        return None, warnings
    return ranked[0], warnings


def load_bos_nre_reference(source_file: Any, configuration: str) -> tuple[str, HolidayNreReference | None, list[str]]:
    selected_configuration = normalize_nre_configuration(configuration)
    payload = _coerce_workbook_bytes(source_file)
    wb = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
    try:
        sheet_name = select_nre_sheet(wb.sheetnames)
        ws = wb[sheet_name]
        titles = _find_nre_titles(ws)
        if selected_configuration not in titles:
            found = ", ".join(sorted(titles)) or "none"
            return sheet_name, None, [f"Could not find title for Holiday NRE {selected_configuration}; found: {found}."]
        images = _extract_reference_images(ws)
        image, warnings = _select_image_near_title(selected_configuration, titles[selected_configuration], images)
        if image is None:
            return sheet_name, None, warnings
        title_text, title_row, title_col = titles[selected_configuration]
        return (
            sheet_name,
            HolidayNreReference(
                configuration=selected_configuration,
                sheet_name=sheet_name,
                title_text=title_text,
                title_row=title_row,
                title_col=title_col,
                image=image,
                warnings=warnings,
            ),
            warnings,
        )
    finally:
        wb.close()


def load_all_bos_nre_references(source_file: Any) -> tuple[str, dict[str, HolidayNreReference], list[str]]:
    references: dict[str, HolidayNreReference] = {}
    warnings: list[str] = []
    sheet_name = ""
    for configuration in NRE_CONFIGURATIONS:
        try:
            sheet_name, reference, config_warnings = load_bos_nre_reference(source_file, configuration)
            warnings.extend(config_warnings)
            if reference is not None:
                references[configuration] = reference
        except Exception as exc:
            warnings.append(f"{configuration}: {exc}")
    return sheet_name, references, warnings
